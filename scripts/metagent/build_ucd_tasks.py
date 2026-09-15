#!/usr/bin/env python3
"""尿素循环障碍(UCD)代谢组 → benchmark 任务(非 TCA、人类、已发表、非循环)。

源:Untargeted metabolomics in UCD, Genet Med 2019(PMID 30670878),mmc3.xlsx。
4 sheet = 4 种酶缺陷,每 sheet = 代谢物 × 患者 z-score(相对对照,z-score 即差异)。
gold = 已知缺陷酶(来自诊断,非循环,不用 SUPER/SUB_PATHWAY 注释)。

用法:python scripts/metagent/build_ucd_tasks.py <mmc3.xlsx> <out_dir>
"""
from __future__ import annotations
import json, sys
from pathlib import Path
import openpyxl

# sheet 名 → (缺陷酶, gold 通路, 诊断 biomarker 关键词, 期望方向)
SUBTYPES = {
    "Final Citrullinemia for Paper": ("ASS1 (argininosuccinate synthase)", "Urea cycle — citrullinemia type I (ASS1)",
                                       ["citrulline"], "up"),
    "Final ASL Def for Paper": ("ASL (argininosuccinate lyase)", "Urea cycle — argininosuccinic aciduria (ASL)",
                                 ["argininosuccinate", "citrulline"], "up"),
    "Final OTCD": ("OTC (ornithine transcarbamylase)", "Urea cycle — OTC deficiency",
                    ["orotate", "orotic", "uracil"], "up"),
    "Final ARG1D": ("ARG1 (arginase 1)", "Urea cycle / arginine — arginase deficiency",
                     ["arginine"], "up"),
}
TOP_N = 15


def main():
    xlsx = Path(sys.argv[1]); out = Path(sys.argv[2]); out.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    tasks = []; log = []
    for sheet, (enzyme, pathway, bm_kw, direction) in SUBTYPES.items():
        if sheet not in wb.sheetnames:
            log.append(f"{sheet}: missing"); continue
        ws = wb[sheet]
        rows = list(ws.iter_rows(values_only=True))
        hdr = list(rows[0])
        # patient z-score columns = after 'Fill Rate'
        try: pstart = hdr.index("Fill Rate") + 1
        except ValueError: pstart = 7
        pcols = [i for i in range(pstart, len(hdr)) if hdr[i] is not None]
        met_i = 0  # BIOCHEMICAL
        diff = []
        for r in rows[1:]:
            name = r[met_i]
            if not name: continue
            zs = []
            for i in pcols:
                try: zs.append(float(r[i]))
                except (TypeError, ValueError): pass
            if len(zs) < 2: continue
            mz = sum(zs) / len(zs)  # mean z-score across patients = differential
            diff.append({"name": str(name), "mean_z": round(mz, 2), "n_pat": len(zs)})
        diff.sort(key=lambda x: -abs(x["mean_z"]))
        top = diff[:TOP_N]
        bm = next((d for d in diff if any(k in d["name"].lower() for k in bm_kw)), None)
        solv = bm is not None and ((bm["mean_z"] > 1 and direction == "up") or (bm["mean_z"] < -1 and direction == "down"))
        tasks.append({
            "task_id": f"ucd_{enzyme.split()[0].lower()}",
            "source": "UCD untargeted metabolomics, Genet Med 2019 (PMID 30670878), mmc3.xlsx",
            "perturbation": {"type": "inborn_error", "disease": sheet.replace("Final ", "").replace(" for Paper", ""),
                             "defective_enzyme": enzyme, "organism": "Homo sapiens", "n_patients": len(pcols)},
            "input": {"differential_metabolites": top},
            "gold": {"defective_enzyme": enzyme, "perturbed_pathway": pathway,
                     "diagnostic_biomarker": bm_kw[0], "gold_anchor": "inborn error (non-circular)",
                     "solvability": "PASS" if solv else "WEAK"},
        })
        log.append(f"{sheet}: n_pat={len(pcols)} | biomarker {bm['name'] if bm else 'NA'} "
                   f"mean_z={bm['mean_z'] if bm else 'NA'} -> {'PASS' if solv else 'WEAK'}")

    (out / "ucd_tasks.jsonl").write_text("\n".join(json.dumps(t, ensure_ascii=False) for t in tasks) + "\n")
    summary = {"n_tasks": len(tasks), "n_pass": sum(1 for t in tasks if t["gold"]["solvability"] == "PASS"), "log": log}
    (out / "ucd_summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2))
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    for t in tasks:
        print(f"\n[{t['task_id']}] gold={t['gold']['perturbed_pathway']}")
        print("  top-5:", [(d['name'][:24], d['mean_z']) for d in t['input']['differential_metabolites'][:5]])


if __name__ == "__main__":
    main()
